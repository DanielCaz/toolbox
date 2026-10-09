"""Read and write tabular data (CSV/TSV, XLSX, JSON/JSONL) as plain lists of rows.

A "table" is ``list[list[Any]]`` whose first row is the header. A "workbook" maps sheet
names to tables, in order. Used by the spreadsheet converter and the PDF table extractor.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.core.errors import ToolError

Table = list[list[Any]]
Workbook_ = dict[str, Table]

MAX_CELLS = 5_000_000
_INT = re.compile(r"^-?(0|[1-9]\d{0,17})$")
_FLOAT = re.compile(r"^-?(0|[1-9]\d*)\.\d+$")
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


# -- helpers -------------------------------------------------------------------
def coerce(value: str) -> Any:
    """Turn a CSV string into int/float/bool/None where that is unambiguous."""
    v = value.strip()
    if v == "":
        return None
    if _INT.match(v):
        return int(v)
    if _FLOAT.match(v):
        return float(v)
    if v.lower() in ("true", "false"):
        return v.lower() == "true"
    return value  # keep strings (including leading-zero ids like 00123) untouched


def sheet_name(name: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("_", name).strip("' ")[:31] or "Sheet"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _trim(rows: Table) -> Table:
    """Drop fully empty trailing rows and columns."""
    while rows and all(c in (None, "") for c in rows[-1]):
        rows.pop()
    if not rows:
        return rows
    width = max(len(r) for r in rows)
    while width > 0 and all(
        (r[width - 1] if len(r) >= width else None) in (None, "") for r in rows
    ):
        width -= 1
    return [list(r[:width]) + [None] * (width - len(r)) for r in rows]


def _check_size(rows: Table, name: str) -> None:
    if sum(len(r) for r in rows) > MAX_CELLS:
        raise ToolError(f"'{name}' is too large to convert (more than {MAX_CELLS:,} cells).")


def header_names(row: list[Any]) -> list[str]:
    """Unique, non-empty header strings (blank → col3, duplicates → name_2)."""
    seen: dict[str, int] = {}
    out: list[str] = []
    for i, cell in enumerate(row, start=1):
        name = str(cell).strip() if cell not in (None, "") else f"col{i}"
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    return out


def _jsonable(value: Any) -> Any:
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return value


# -- detect + read ---------------------------------------------------------------
def detect_kind(path: Path, mime: str) -> str:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return "xlsx"
    if ext == ".json":
        return "json"
    if ext in (".jsonl", ".ndjson"):
        return "jsonl"
    if ext == ".tsv":
        return "tsv"
    if ext == ".csv":
        return "csv"
    if "spreadsheetml" in mime:
        return "xlsx"
    if mime == "application/json":
        return "json"
    return "csv"


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_delimited(path: Path, kind: str, infer_types: bool) -> Table:
    text = _decode(path.read_bytes())
    if kind == "tsv":
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    try:
        rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    except csv.Error as e:
        raise ToolError(f"'{path.name}' is not valid CSV: {e}") from None
    if infer_types and rows:
        rows = [rows[0]] + [[coerce(c) for c in r] for r in rows[1:]]
    return _trim(rows)


def read_xlsx(path: Path, only_sheet: str = "") -> Workbook_:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        raise ToolError(f"'{path.name}' could not be opened as an .xlsx workbook.") from None
    try:
        names = wb.sheetnames
        if only_sheet:
            if only_sheet not in names:
                raise ToolError(f"No sheet called '{only_sheet}'. Sheets: {', '.join(names)}.")
            names = [only_sheet]
        result: Workbook_ = {}
        for name in names:
            rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
            _check_size(rows, path.name)
            result[name] = _trim(rows)
        return result
    finally:
        wb.close()


def _records_to_table(records: list[Any], name: str) -> Table:
    if all(isinstance(r, dict) for r in records):
        keys: list[str] = []
        for r in records:
            for k in r:
                if k not in keys:
                    keys.append(k)
        return [keys] + [[_flat(r.get(k)) for k in keys] for r in records]
    if all(isinstance(r, list) for r in records):
        return [[_flat(c) for c in r] for r in records]
    raise ToolError(f"'{name}': expected a list of objects or a list of rows.")


def _flat(value: Any) -> Any:
    """Nested JSON becomes a JSON string so it fits in a cell."""
    return json.dumps(value, ensure_ascii=False) if isinstance(value, dict | list) else value


def read_json(path: Path, kind: str) -> Workbook_:
    text = _decode(path.read_bytes())
    try:
        if kind == "jsonl":
            data: Any = [json.loads(line) for line in text.splitlines() if line.strip()]
        else:
            data = json.loads(text)
    except json.JSONDecodeError as e:
        raise ToolError(f"'{path.name}' is not valid JSON: {e.msg} (line {e.lineno}).") from None

    if isinstance(data, list):
        return {"Sheet1": _trim(_records_to_table(data, path.name))}
    if isinstance(data, dict):
        if data and all(isinstance(v, list) for v in data.values()):  # {sheet: [records]}
            return {k: _trim(_records_to_table(v, path.name)) for k, v in data.items()}
        return {"Sheet1": _trim(_records_to_table([data], path.name))}
    raise ToolError(f"'{path.name}': the JSON must be an object or a list.")


def read_any(path: Path, kind: str, *, infer_types: bool = True, sheet: str = "") -> Workbook_:
    if kind == "xlsx":
        return read_xlsx(path, sheet)
    if kind in ("json", "jsonl"):
        return read_json(path, kind)
    return {"Sheet1": read_delimited(path, kind, infer_types)}


# -- write -----------------------------------------------------------------------
def write_delimited(table: Table, path: Path, *, tab: bool = False, bom: bool = True) -> None:
    with path.open("w", newline="", encoding="utf-8-sig" if bom else "utf-8") as fh:
        writer = csv.writer(fh, delimiter="\t" if tab else ",")
        for row in table:
            writer.writerow(
                [
                    "" if c is None else str(c).lower() if isinstance(c, bool) else _jsonable(c)
                    for c in row
                ]
            )


def table_records(table: Table) -> list[dict[str, Any]]:
    if not table:
        return []
    names = header_names(table[0])
    return [
        {n: _jsonable(r[i] if i < len(r) else None) for i, n in enumerate(names)} for r in table[1:]
    ]


def write_json(book: Workbook_, path: Path) -> None:
    if len(book) == 1:
        payload: Any = table_records(next(iter(book.values())))
    else:
        payload = {name: table_records(t) for name, t in book.items()}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_jsonl(table: Table, path: Path) -> None:
    with path.open("w", encoding="utf-8") as fh:
        for rec in table_records(table):
            fh.write(json.dumps(rec, ensure_ascii=False) + "\n")


def write_xlsx(book: Workbook_, path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for name, table in book.items():
        ws = wb.create_sheet(sheet_name(name, used))
        widths: dict[int, int] = {}
        for r_idx, row in enumerate(table, start=1):
            for c_idx, cell in enumerate(row, start=1):
                if cell is None or cell == "":
                    continue
                ws.cell(row=r_idx, column=c_idx, value=cell)
                widths[c_idx] = max(widths.get(c_idx, 0), len(str(cell)))
        for c_idx, w in widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = min(60, max(8, w + 2))
        if table:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            ws.freeze_panes = "A2"
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(path)
