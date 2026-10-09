import io
import tarfile
import zipfile

import py7zr
import pytest
import zxingcpp
from openpyxl import load_workbook
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

from app.core.errors import ToolError
from app.tools.files import archive as archive_mod
from app.tools.files.archive import Archive, flat_name
from app.tools.ocr.tables import PdfTables
from app.tools.qr.read import ReadQr


# ------------------------------------------------------------------ PDF tables
@pytest.fixture
def table_pdf(tmp_path):
    path = tmp_path / "inputs" / "report.pdf"
    path.parent.mkdir(exist_ok=True)
    data = [["Region", "Q1", "Q2"], ["EU", "10", "12"], ["US", "20", "25"], ["APAC", "5", "9"]]
    t = Table(data)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(str(path), pagesize=A4).build([t])
    return path


def test_pdf_tables_to_xlsx(table_pdf, run_tool):
    (out,) = run_tool(PdfTables, [table_pdf], output="xlsx")
    ws = load_workbook(out).worksheets[0]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["Region", "Q1", "Q2"]
    assert rows[2] == ["US", "20", "25"]
    assert out.name == "report-tables.xlsx"


def test_pdf_tables_to_csv(table_pdf, run_tool):
    (out,) = run_tool(PdfTables, [table_pdf], output="csv")
    assert out.name == "report-p1-t1.csv"
    assert "EU,10,12" in out.read_text(encoding="utf-8-sig")


def test_pdf_without_tables_explains_what_to_try(make_pdf, run_tool):
    with pytest.raises(ToolError, match="text"):
        run_tool(PdfTables, [make_pdf("plain.pdf", 1)])
    with pytest.raises(ToolError, match="OCR"):
        run_tool(PdfTables, [make_pdf("plain.pdf", 1)], strategy="text")


# --------------------------------------------------------------------- archives
def make_zip(path, members: dict[str, bytes]):
    path.parent.mkdir(exist_ok=True)
    with zipfile.ZipFile(path, "w") as z:
        for name, data in members.items():
            z.writestr(name, data)
    return path


def test_flat_name_never_returns_a_path():
    assert flat_name("docs/sub/readme.txt") == "docs_sub_readme.txt"
    assert flat_name("../../etc/passwd") == "etc_passwd"
    assert flat_name("/abs/olute.txt") == "abs_olute.txt"
    assert flat_name("..\\..\\win.ini") == "win.ini"
    assert flat_name("dir/") == "dir"
    assert flat_name("../") == ""


@pytest.mark.parametrize("fmt", ["zip", "tar.gz", "7z"])
def test_create_then_extract_round_trip(tmp_path, run_tool, fmt):
    a = tmp_path / "inputs" / "a.txt"
    a.parent.mkdir(exist_ok=True)
    a.write_text("alpha")
    b = tmp_path / "inputs" / "b.txt"
    b.write_text("beta")
    (arc,) = run_tool(Archive, [a, b], action="create", format=fmt, output_name="bundle")
    assert arc.name == f"bundle.{fmt}"

    outs = run_tool(Archive, [arc], action="extract")
    assert {o.name: o.read_text() for o in outs} == {"a.txt": "alpha", "b.txt": "beta"}


def test_extract_flattens_nested_and_hostile_names(tmp_path, run_tool):
    src = make_zip(
        tmp_path / "inputs" / "x.zip",
        {"docs/a.txt": b"A", "../../evil.txt": b"E", "/abs.txt": b"B", "empty_dir/": b""},
    )
    outs = run_tool(Archive, [src], action="extract")
    assert sorted(o.name for o in outs) == ["abs.txt", "docs_a.txt", "evil.txt"]
    for o in outs:  # everything stays inside the job's output folder
        assert o.parent.name == "out"


def test_extract_skips_symlinks_in_tar(tmp_path, run_tool):
    path = tmp_path / "inputs" / "l.tar"
    path.parent.mkdir(exist_ok=True)
    with tarfile.open(path, "w") as tf:
        info = tarfile.TarInfo("real.txt")
        info.size = 2
        tf.addfile(info, io.BytesIO(b"ok"))
        link = tarfile.TarInfo("shadow")
        link.type = tarfile.SYMTYPE
        link.linkname = "/etc/shadow"
        tf.addfile(link)
    outs = run_tool(Archive, [path], action="extract")
    assert [o.name for o in outs] == ["real.txt"]


def test_extract_limits(tmp_path, run_tool, monkeypatch):
    src = make_zip(tmp_path / "inputs" / "big.zip", {"a.bin": b"x" * 5000, "b.bin": b"y" * 5000})
    monkeypatch.setattr(archive_mod, "MAX_TOTAL_BYTES", 6000)
    with pytest.raises(ToolError, match="2 GB"):
        run_tool(Archive, [src], action="extract")
    monkeypatch.setattr(archive_mod, "MAX_TOTAL_BYTES", 10**9)
    monkeypatch.setattr(archive_mod, "MAX_MEMBERS", 1)
    with pytest.raises(ToolError, match="more than 1 files"):
        run_tool(Archive, [src], action="extract")


def test_extract_errors(tmp_path, run_tool):
    junk = tmp_path / "inputs" / "junk.bin"
    junk.parent.mkdir(exist_ok=True)
    junk.write_bytes(b"not an archive at all")
    with pytest.raises(ToolError, match="not a zip"):
        run_tool(Archive, [junk], action="extract")
    with pytest.raises(ToolError, match="exactly one"):
        run_tool(Archive, [junk, junk], action="extract")

    empty = make_zip(tmp_path / "inputs" / "empty.zip", {})
    with pytest.raises(ToolError, match="empty"):
        run_tool(Archive, [empty], action="extract")

    enc = tmp_path / "inputs" / "locked.7z"
    with py7zr.SevenZipFile(enc, "w", password="pw") as sz:
        sz.writestr(b"secret", "s.txt")
    with pytest.raises(ToolError, match="password"):
        run_tool(Archive, [enc], action="extract")


# ------------------------------------------------------------------- QR reading
def qr_image(tmp_path, text, name="qr.png", fmt=zxingcpp.BarcodeFormat.QRCode):
    bc = zxingcpp.create_barcode(text, fmt)
    img = Image.fromarray(bc.to_image(scale=8)).convert("RGB")
    path = tmp_path / "inputs" / name
    path.parent.mkdir(exist_ok=True)
    img.save(path)
    return path


def test_read_qr_code(tmp_path, run_tool):
    text = run_tool(ReadQr, [qr_image(tmp_path, "https://example.com/hello?x=1")])
    assert text == "[QRCode] https://example.com/hello?x=1"


def test_read_barcode_and_multiple_files(tmp_path, run_tool):
    a = qr_image(tmp_path, "first")
    b = qr_image(tmp_path, "5901234123457", "ean.png", zxingcpp.BarcodeFormat.EAN13)
    text = run_tool(ReadQr, [a, b])
    assert "=== qr.png ===\n[QRCode] first" in text
    assert "[EAN13] 5901234123457" in text


def test_read_qr_nothing_found(make_image, run_tool):
    with pytest.raises(ToolError, match="No QR code"):
        run_tool(ReadQr, [make_image("blank.png")])
