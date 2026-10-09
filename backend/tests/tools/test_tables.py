import json

import pytest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader  # noqa: F401  (keeps import style consistent)

from app.core import tables
from app.core.errors import ToolError
from app.tools.docs.tabular import Tabular


def write(tmp_path, name, content, mode="w", **kw):
    p = tmp_path / "inputs" / name
    p.parent.mkdir(exist_ok=True)
    if isinstance(content, bytes):
        p.write_bytes(content)
    else:
        p.write_text(content, **kw)
    return p


def test_csv_to_xlsx_infers_types_but_keeps_leading_zeros(tmp_path, run_tool):
    src = write(
        tmp_path, "people.csv", "id,name,age,score,active\n00123,Ann,31,4.5,true\n7,Bob,,2,false\n"
    )
    (out,) = run_tool(Tabular, [src], to="xlsx")
    ws = load_workbook(out).active
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    assert rows[0] == ["id", "name", "age", "score", "active"]
    assert rows[1] == ["00123", "Ann", 31, 4.5, True]
    assert rows[2] == [7, "Bob", None, 2, False]
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def test_csv_without_type_inference_keeps_text(tmp_path, run_tool):
    src = write(tmp_path, "t.csv", "a,b\n1,2\n")
    (out,) = run_tool(Tabular, [src], to="json", infer_types=False)
    assert json.loads(out.read_text()) == [{"a": "1", "b": "2"}]


def test_delimiter_and_encoding_detection(tmp_path, run_tool):
    semi = write(tmp_path, "eu.csv", "name;city\nJosé;Zürich\n".encode("cp1252"))
    (out,) = run_tool(Tabular, [semi], to="json")
    assert json.loads(out.read_text(encoding="utf-8")) == [{"name": "José", "city": "Zürich"}]


def test_json_to_csv_flattens_nested_and_unions_keys(tmp_path, run_tool):
    data = [{"a": 1, "tags": ["x", "y"]}, {"a": 2, "b": {"k": 1}}]
    src = write(tmp_path, "d.json", json.dumps(data))
    (out,) = run_tool(Tabular, [src], to="csv", excel_bom=False)
    lines = out.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "a,tags,b"
    assert lines[1] == '1,"[""x"", ""y""]",'
    assert lines[2] == '2,,"{""k"": 1}"'


def test_csv_output_has_excel_bom_by_default(tmp_path, run_tool):
    src = write(tmp_path, "d.json", '[{"n": "é"}]')
    (out,) = run_tool(Tabular, [src], to="csv")
    assert out.read_bytes().startswith(b"\xef\xbb\xbf")


def test_jsonl_round_trip(tmp_path, run_tool):
    src = write(tmp_path, "e.csv", "k,v\na,1\nb,2\n")
    (jl,) = run_tool(Tabular, [src], to="jsonl")
    assert [json.loads(line) for line in jl.read_text().splitlines()] == [
        {"k": "a", "v": 1},
        {"k": "b", "v": 2},
    ]
    (back,) = run_tool(Tabular, [jl], to="tsv", excel_bom=False)
    assert back.read_text().splitlines() == ["k\tv", "a\t1", "b\t2"]


def test_multi_sheet_xlsx_to_csv_gives_one_file_per_sheet(tmp_path, run_tool):
    wb = Workbook()
    wb.active.title = "Sales"
    wb.active.append(["q", "total"])
    wb.active.append(["Q1", 10])
    wb.create_sheet("Costs").append(["item", "eur"])
    wb["Costs"].append(["rent", 5])
    path = tmp_path / "inputs" / "book.xlsx"
    path.parent.mkdir(exist_ok=True)
    wb.save(path)

    outs = run_tool(Tabular, [path], to="csv", excel_bom=False)
    assert sorted(o.name for o in outs) == ["book-Costs.csv", "book-Sales.csv"]
    only = run_tool(Tabular, [path], to="json", sheet="Costs")
    assert json.loads(only[0].read_text()) == [{"item": "rent", "eur": 5}]
    both = run_tool(Tabular, [path], to="json")
    assert set(json.loads(both[0].read_text())) == {"Sales", "Costs"}
    with pytest.raises(ToolError, match="No sheet called"):
        run_tool(Tabular, [path], to="csv", sheet="Nope")


def test_errors(tmp_path, run_tool):
    with pytest.raises(ToolError, match="already CSV"):
        run_tool(Tabular, [write(tmp_path, "a.csv", "x\n1\n")], to="csv")
    with pytest.raises(ToolError, match="not valid JSON"):
        run_tool(Tabular, [write(tmp_path, "bad.json", "{oops")], to="csv")
    with pytest.raises(ToolError, match="no data"):
        run_tool(Tabular, [write(tmp_path, "empty.csv", "\n\n")], to="json")
    with pytest.raises(ToolError, match=".xlsx"):
        run_tool(Tabular, [write(tmp_path, "fake.xlsx", "not a workbook")], to="csv")


def test_sheet_name_sanitising():
    used: set[str] = set()
    assert tables.sheet_name("a/b:c", used) == "a_b_c"
    assert tables.sheet_name("a/b:c", used) == "a_b_c (2)"
    assert len(tables.sheet_name("x" * 50, set())) == 31
